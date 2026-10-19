# make_test_docx.py
"""
Генератор тестового документа: таблица переносится на несколько страниц.
Запуск: python make_test_docx.py
Результат: test_natural_break.docx
"""
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUTPUT = 'test_natural_break.docx'


def page_break_in(cell):
    """Явный разрыв страницы внутри ячейки."""
    run = OxmlElement('w:r')
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    run.append(br)
    cell.paragraphs[0]._element.append(run)


def main():
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)

    # --- Заголовок ---
    doc.add_paragraph('1. ЭКСПЕРИМЕНТАЛЬНАЯ ЧАСТЬ')

    # --- Много текста перед таблицей, чтобы она начиналась в конце страницы ---
    for i in range(15):
        doc.add_paragraph(
            f'Это абзац текста номер {i + 1}. Он занимает место на странице, '
            'чтобы таблица началась в самом низу и туда влезло только 2-3 строки. '
            'Остальные строки должны перенестись на следующую страницу.'
        )

    # --- Таблица ---
    p = doc.add_paragraph('Таблица 1 — Результаты измерений')
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0)

    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    
    hdr = table.rows[0].cells
    hdr[0].text = '№'
    hdr[1].text = 'Параметр'
    hdr[2].text = 'Значение'

    # Добавляем 50 строк данных
    for i in range(1, 51):
        cells = table.add_row().cells
        cells[0].text = str(i)
        cells[1].text = f'Параметр {i}'
        cells[2].text = f'{i * 1.5:.2f}'

        # Имитируем естественный перенос:
        # - После 2-й строки (таблица началась в конце страницы, влезло только 2 строки)
        # - Потом каждые 25 строк (остальные переносятся на следующие страницы)
        if i in (2, 27):
            page_break_in(cells[0])

    # --- Текст после таблицы ---
    doc.add_paragraph(
        'После длинной таблицы идёт обычный абзац текста. '
        'Он должен остаться между таблицей и следующим блоком.'
    )

    # --- Короткая таблица (не должна разбиваться) ---
    p2 = doc.add_paragraph('Таблица 2 — Короткая таблица')
    p2.paragraph_format.first_line_indent = Cm(0)
    p2.paragraph_format.left_indent = Cm(0)

    table2 = doc.add_table(rows=1, cols=2)
    table2.style = 'Table Grid'
    
    hdr2 = table2.rows[0].cells
    hdr2[0].text = '№'
    hdr2[1].text = 'Значение'

    for i in range(1, 6):
        cells = table2.add_row().cells
        cells[0].text = str(i)
        cells[1].text = f'{i * 10}'

    doc.save(OUTPUT)

    print(f'Готово: {OUTPUT}')
    print('Таблица 1: 50 строк, разрывы после строк 2 и 27')
    print('Таблица 2: 5 строк (не должна разбиваться)')
    print('Загрузи этот файл на сайт и проверь результат.')


if __name__ == '__main__':
    main()