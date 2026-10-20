from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL, WD_LINE_SPACING as LS
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
import re
import copy

TNR = 'Times New Roman'
BLACK = '000000'

RE_H3 = re.compile(r'^\d+\.\d+\.\d+(?:[.)]\s*|\s+)(?=\S)')
RE_H2 = re.compile(r'^\d+\.\d+(?:[.)]\s*|\s+)(?=\S)')
RE_H1 = re.compile(r'^\d+(?:[.)]\s*|\s+)(?=\S)')
RE_TOC = re.compile(r'[.…]{2,}\s*\d+\s*$')


# ==================== XML-УТИЛИТЫ ====================

def _v(tag, val):
    el = OxmlElement(f'w:{tag}')
    el.set(qn('w:val'), str(val))
    return el


def _local(el):
    return el.tag.rsplit('}', 1)[-1]


def _drop(parent, *tags):
    for tag in tags:
        for el in parent.findall(qn(f'w:{tag}')):
            parent.remove(el)


def _lang():
    el = OxmlElement('w:lang')
    el.set(qn('w:val'), 'ru-RU')
    el.set(qn('w:eastAsia'), 'ru-RU')
    return el


# ==================== ШРИФТ ====================

def _format_run(r, size=14, bold=False, italic=False):
    old = r.find(qn('w:rPr'))
    keep = []
    has_lang = False

    if old is not None:
        for child in list(old):
            tag = _local(child)
            if tag in ('vertAlign', 'rtl'):
                keep.append(child)
            if tag == 'lang':
                has_lang = True
        r.remove(old)

    rPr = OxmlElement('w:rPr')

    rf = OxmlElement('w:rFonts')
    for attr in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
        rf.set(qn(f'w:{attr}'), TNR)
    rPr.append(rf)

    rPr.append(_v('b', 1 if bold else 0))
    rPr.append(_v('bCs', 1 if bold else 0))
    rPr.append(_v('i', 1 if italic else 0))
    rPr.append(_v('iCs', 1 if italic else 0))
    rPr.append(_v('color', BLACK))
    rPr.append(_v('sz', int(size * 2)))
    rPr.append(_v('szCs', int(size * 2)))

    for el in keep:
        rPr.append(el)
    if not has_lang:
        rPr.append(_lang())

    r.insert(0, rPr)


def style_runs(p, size=14, bold=False, italic=False):
    for r in p._element.iter(qn('w:r')):
        _format_run(r, size, bold, italic)


def upper_text(p):
    for t in p._element.iter(qn('w:t')):
        if t.text:
            t.text = t.text.upper()


def replace_text(p, new_text):
    runs = list(p._element.iter(qn('w:r')))
    if not runs:
        return
    first_t = None
    for r in runs:
        for t in r.findall(qn('w:t')):
            if first_t is None:
                first_t = t
            else:
                t.getparent().remove(t)
    if first_t is not None:
        first_t.text = new_text
        first_t.set(qn('xml:space'), 'preserve')


# ==================== ПАРАГРАФ ====================

def clear_p(p):
    pPr = p._element.get_or_add_pPr()
    for child in list(pPr):
        if _local(child) not in ('pStyle', 'numPr', 'sectPr'):
            pPr.remove(child)


def fmt_p(p, align=AL.JUSTIFY, indent=1.25, left=0.0, line=1.5,
          before=0, after=0, outline=None, keep_next=False):
    pf = p.paragraph_format
    pf.alignment = align
    pf.line_spacing_rule = LS.MULTIPLE
    pf.line_spacing = line
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.first_line_indent = Cm(indent)
    pf.left_indent = Cm(left)
    pf.right_indent = Cm(0)
    pf.keep_with_next = keep_next

    pPr = p._element.get_or_add_pPr()
    _drop(pPr, 'outlineLvl')

    if outline is not None:
        el = _v('outlineLvl', outline)
        for child in pPr:
            if _local(child) in ('divId', 'cnfStyle', 'rPr', 'sectPr', 'pPrChange'):
                child.addprevious(el)
                break
        else:
            pPr.append(el)


# ==================== СТИЛИ ДОКУМЕНТА ====================

def _heading_style(doc, level):
    targets = {f'heading{level}', f'заголовок{level}'}
    for s in doc.styles:
        if s.type != WD_STYLE_TYPE.PARAGRAPH:
            continue
        name = (s.name or '').lower().replace(' ', '').replace('\xa0', '')
        if name in targets:
            return s
    return doc.styles.add_style(f'Heading {level}', WD_STYLE_TYPE.PARAGRAPH)


def _reset_style(s, *, size, bold, align, indent, outline, normal_id,
                 is_normal=False):
    e = s._element

    name_el = e.find(qn('w:name'))
    name_val = (name_el.get(qn('w:val')) if name_el is not None
                else (s.name or 'Normal'))

    for child in list(e):
        e.remove(child)

    nm = OxmlElement('w:name')
    nm.set(qn('w:val'), name_val)
    e.append(nm)

    if not is_normal:
        e.append(_v('basedOn', normal_id))
        e.append(_v('next', normal_id))
    e.append(OxmlElement('w:qFormat'))

    pf = s.paragraph_format
    pf.alignment = align
    pf.line_spacing_rule = LS.MULTIPLE
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Cm(indent)
    pf.left_indent = Cm(0)
    pf.right_indent = Cm(0)
    pf.keep_with_next = outline is not None

    pPr = e.find(qn('w:pPr'))
    if outline is not None and pPr is not None:
        pPr.append(_v('outlineLvl', outline))

    rPr = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts')
    for attr in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
        rf.set(qn(f'w:{attr}'), TNR)
    rPr.append(rf)
    rPr.append(_v('b', 1 if bold else 0))
    rPr.append(_v('bCs', 1 if bold else 0))
    rPr.append(_v('color', BLACK))
    rPr.append(_v('sz', int(size * 2)))
    rPr.append(_v('szCs', int(size * 2)))
    rPr.append(_lang())
    e.append(rPr)


def setup_styles(doc):
    normal = doc.styles['Normal']
    normal_id = normal._element.get(qn('w:styleId'))

    _reset_style(normal, size=14, bold=False, align=AL.JUSTIFY,
                 indent=1.25, outline=None, normal_id=normal_id,
                 is_normal=True)

    for level in (1, 2, 3):
        s = _heading_style(doc, level)
        _reset_style(
            s, size=14, bold=True,
            align=AL.CENTER if level == 1 else AL.LEFT,
            indent=0.0 if level == 1 else 1.25,
            outline=level - 1,
            normal_id=normal_id,
        )
        if level == 1:
            s.paragraph_format.line_spacing = 2.0


# ==================== РАЗРЫВЫ СТРАНИЦ ====================

def ensure_page_break(p):
    prev = p._element.getprevious()
    while prev is not None and prev.tag != qn('w:p'):
        prev = prev.getprevious()

    if prev is None:
        return

    for br in prev.iter(qn('w:br')):
        if br.get(qn('w:type')) == 'page':
            return

    pPr = prev.find(qn('w:pPr'))
    if pPr is not None and pPr.find(qn('w:sectPr')) is not None:
        return

    br_run = OxmlElement('w:r')
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    br_run.append(br)
    prev.append(br_run)


# ==================== РАСПОЗНАВАНИЕ ====================

def _style_heading_level(p, by_id):
    pPr = p._element.pPr
    if pPr is not None:
        el = pPr.find(qn('w:outlineLvl'))
        if el is not None:
            try:
                v = int(el.get(qn('w:val')))
                if v < 3:
                    return v + 1
            except (TypeError, ValueError):
                pass

    s = p.style
    seen = set()
    while s is not None:
        sid = s._element.get(qn('w:styleId'))
        if sid in seen:
            break
        seen.add(sid)

        name = (s.name or '').lower().replace(' ', '')
        m = re.fullmatch(r'(?:heading|заголовок)([1-3])', name)
        if m:
            return int(m.group(1))

        m = re.fullmatch(r'heading([1-3])', (sid or '').lower())
        if m:
            return int(m.group(1))

        sp = s._element.find(qn('w:pPr'))
        if sp is not None:
            el = sp.find(qn('w:outlineLvl'))
            if el is not None:
                try:
                    v = int(el.get(qn('w:val')))
                    if v < 3:
                        return v + 1
                except (TypeError, ValueError):
                    pass

        base = s._element.find(qn('w:basedOn'))
        if base is None:
            break
        s = by_id.get(base.get(qn('w:val')))

    return None


def _looks_like_heading(text):
    if len(text) > 150:
        return False
    body = re.sub(r'^\d+(?:\.\d+)*[.)]?\s*', '', text)
    return bool(body) and body[0].isupper()


def detect(p, in_bib, by_id):
    text = p.text.strip()
    if not text:
        return 'empty'
    up = text.upper()

    if RE_TOC.search(text):
        return 'toc_entry'

    if up in ('СОДЕРЖАНИЕ', 'ОГЛАВЛЕНИЕ'):
        return 'contents_title'

    if up in ('ВВЕДЕНИЕ', 'ВВОДНАЯ ЧАСТЬ'):
        return 'heading1'

    if up in ('ЗАКЛЮЧЕНИЕ', 'ВЫВОДЫ'):
        return 'heading1'

    if up.startswith('СПИСОК') and ('ИСТОЧНИК' in up or 'ЛИТЕРАТУР' in up):
        return 'bib_title'
    if up in ('ЛИТЕРАТУРА', 'БИБЛИОГРАФИЯ', 'БИБЛИОГРАФИЧЕСКИЙ СПИСОК',
              'ИСПОЛЬЗОВАННЫЕ ИСТОЧНИКИ', 'ПЕРЕЧЕНЬ ССЫЛОК'):
        return 'bib_title'

    if up.startswith('ПРИЛОЖЕНИЕ'):
        return 'appendix'

    if in_bib and re.match(r'^(\[\d+\]|\d+[\.)])', text):
        return 'bib_item'

    if re.match(r'^(рисунок|рис\.)\s*\d+', text, re.I):
        return 'fig_caption'
    if re.match(r'^(таблица|таб\.)\s*\d+', text, re.I):
        return 'table_caption'

    if not in_bib and _looks_like_heading(text):
        if RE_H3.match(text):
            return 'heading3'
        if RE_H2.match(text):
            return 'heading2'
        if RE_H1.match(text):
            return 'heading1'

    lvl = _style_heading_level(p, by_id)
    if lvl and len(text) <= 150:
        return f'heading{lvl}'

    if re.match(r'^ГЛАВА\s+\d+', up):
        return 'heading1'

    if (up == text and any(c.isalpha() for c in text)
            and len(text) <= 120 and p.alignment == AL.CENTER):
        return 'heading1'

    return 'normal'


# ==================== ОБРАБОТЧИКИ ====================

def apply_heading(p, level, hs):
    p.style = hs[level]
    style_runs(p, size=14, bold=True)
    if level == 1:
        upper_text(p)
    fmt_p(p,
          align=AL.CENTER if level == 1 else AL.LEFT,
          indent=0.0 if level == 1 else 1.25,
          line=2.0 if level == 1 else 1.5,
          outline=level - 1,
          keep_next=True)


def apply_body(p, doc, size=14, align=AL.JUSTIFY, indent=1.25, line=1.5):
    p.style = doc.styles['Normal']
    style_runs(p, size=size, bold=False)
    fmt_p(p, align=align, indent=indent, line=line)


def insert_blank_line_after_tables(doc):
    """Вставляет пустой параграф после таблицы, НО НЕ после таблиц-продолжений."""
    body = doc.element.body
    tables = body.findall(qn('w:tbl'))

    for tbl in tables:
        # Проверяем, является ли таблица продолжением
        prev_el = tbl.getprevious()
        is_continuation = False
        if prev_el is not None and prev_el.tag == qn('w:p'):
            p_text = prev_el.text or ''
            if re.match(r'^(продолжение|окончание)\s+таблицы\s+\d+', p_text, re.I):
                is_continuation = True

        if is_continuation:
            continue

        # Проверяем, есть ли уже пустая строка
        next_el = tbl.getnext()
        if next_el is not None and next_el.tag == qn('w:p'):
            next_p = Paragraph(next_el, doc)
            if not next_p.text.strip():
                continue

        new_p = OxmlElement('w:p')
        tbl.addnext(new_p)
        blank = Paragraph(new_p, doc)
        blank.style = doc.styles['Normal']
        clear_p(blank)
        fmt_p(blank, align=AL.LEFT, indent=0, left=0, line=1.0, before=0, after=0)


# ==================== ТАБЛИЦЫ ====================

def format_tables(doc):
    for table in doc.tables:
        for i, row in enumerate(table.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    clear_p(p)
                    p.style = doc.styles['Normal']
                    header = (i == 0)
                    style_runs(p, size=12, bold=header)
                    fmt_p(p,
                          align=AL.CENTER if header else AL.JUSTIFY,
                          indent=0, left=0, line=1.0)


def insert_blank_line_after_tables(doc):
    """Вставляет пустой параграф после таблицы, НО НЕ после таблиц-продолжений."""
    body = doc.element.body
    tables = body.findall(qn('w:tbl'))

    for tbl in tables:
        # Проверяем, является ли таблица продолжением
        prev_el = tbl.getprevious()
        is_continuation = False
        if prev_el is not None and prev_el.tag == qn('w:p'):
            p_text = prev_el.text or ''
            if re.match(r'^(продолжение|окончание)\s+таблицы\s+\d+', p_text, re.I):
                is_continuation = True

        if is_continuation:
            continue

        # Проверяем, есть ли уже пустая строка
        next_el = tbl.getnext()
        if next_el is not None and next_el.tag == qn('w:p'):
            next_p = Paragraph(next_el, doc)
            if not next_p.text.strip():
                continue

        new_p = OxmlElement('w:p')
        tbl.addnext(new_p)
        blank = Paragraph(new_p, doc)
        blank.style = doc.styles['Normal']
        clear_p(blank)
        fmt_p(blank, align=AL.LEFT, indent=0, left=0, line=1.0, before=0, after=0)


def split_long_tables(doc):
    """Разбивает таблицы по явным разрывам страниц в ячейках.
    Шапка повторяется в каждой части."""
    body = doc.element.body
    tables = body.findall(qn('w:tbl'))

    for tbl in tables:
        prev_el = tbl.getprevious()
        table_title = None
        while prev_el is not None:
            if prev_el.tag == qn('w:p'):
                p = Paragraph(prev_el, doc)
                if re.match(r'^(таблица|таб\.)\s*\d+', p.text, re.I):
                    table_title = p.text
                    break
            prev_el = prev_el.getprevious()

        if not table_title:
            continue

        m = re.search(r'(\d+)', table_title)
        if not m:
            continue
        tbl_num = m.group(1)

        rows = tbl.findall(qn('w:tr'))
        if len(rows) <= 1:
            continue

        # Ищем разрывы страниц В ЯЧЕЙКАХ строк (пропускаем шапку — индекс 0)
        page_breaks = []
        for i in range(1, len(rows)):
            row = rows[i]
            has_break = False
            for cell in row.findall(qn('w:tc')):
                for br in cell.iter(qn('w:br')):
                    if br.get(qn('w:type')) == 'page':
                        has_break = True
                        break
                if has_break:
                    break
            if has_break:
                page_breaks.append(i)

        if not page_breaks:
            continue

        # Выделяем шапку и данные
        header = rows[0]
        data_rows = rows[1:]

        # Переводим индексы из rows в data_rows
        data_breaks = [br - 1 for br in page_breaks]

        # Разбиваем данные на части
        chunks = []
        start = 0
        for br_idx in data_breaks:
            chunks.append(data_rows[start:br_idx])
            start = br_idx
        chunks.append(data_rows[start:])

        if len(chunks) <= 1:
            continue

        tblPr = tbl.find(qn('w:tblPr'))
        tblGrid = tbl.find(qn('w:tblGrid'))

        # Первая часть: шапка + chunks[0]
        keep_rows = [header] + chunks[0]
        keep_ids = set(id(r) for r in keep_rows)
        for row in rows:
            if id(row) not in keep_ids:
                tbl.remove(row)

        # Удаляем разрывы страниц из первой части (они больше не нужны)
        for row in keep_rows:
            for cell in row.findall(qn('w:tc')):
                for br in list(cell.iter(qn('w:br'))):
                    if br.get(qn('w:type')) == 'page':
                        br.getparent().remove(br)

        # Остальные части: шапка + chunk
        current_tbl = tbl
        for chunk_idx in range(1, len(chunks)):
            chunk = chunks[chunk_idx]
            is_last = (chunk_idx == len(chunks) - 1)
            caption_text = f"{'Окончание' if is_last else 'Продолжение'} таблицы {tbl_num}"

            # Подпись
            caption_p = OxmlElement('w:p')
            p = Paragraph(caption_p, doc)
            p.style = doc.styles['Normal']
            run = p.add_run(caption_text)
            run.font.name = TNR
            run.font.size = Pt(12)
            p.alignment = AL.RIGHT
            clear_p(p)
            fmt_p(p, align=AL.RIGHT, indent=0, left=0, line=1.0)

            # Новая таблица: ШАПКА + chunk
            new_tbl = OxmlElement('w:tbl')
            if tblPr is not None:
                new_tbl.append(copy.deepcopy(tblPr))
            if tblGrid is not None:
                new_tbl.append(copy.deepcopy(tblGrid))

            new_tbl.append(copy.deepcopy(header))  # ШАПКА КОПИРУЕТСЯ
            for row in chunk:
                new_tbl.append(copy.deepcopy(row))

            # Удаляем разрывы страниц из новой таблицы
            for row_elem in new_tbl.iter(qn('w:tr')):
                for cell in row_elem.findall(qn('w:tc')):
                    for br in list(cell.iter(qn('w:br'))):
                        if br.get(qn('w:type')) == 'page':
                            br.getparent().remove(br)

            current_tbl.addnext(new_tbl)
            new_tbl.addprevious(caption_p)
            current_tbl = new_tbl


# ==================== НОМЕРА СТРАНИЦ ====================

def _field_paragraph(p, instruction, placeholder):
    def add_run(child):
        r = OxmlElement('w:r')
        _format_run(r, size=11)
        r.append(child)
        p._element.append(r)

    def fld(kind):
        el = OxmlElement('w:fldChar')
        el.set(qn('w:fldCharType'), kind)
        return el

    add_run(fld('begin'))

    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = instruction
    add_run(instr)

    add_run(fld('separate'))

    r = OxmlElement('w:r')
    _format_run(r, size=11)
    t = OxmlElement('w:t')
    t.text = placeholder
    r.append(t)
    p._element.append(r)

    add_run(fld('end'))


def add_page_numbers(doc):
    for section in doc.sections:
        section.different_first_page_header_footer = True

        first = section.first_page_footer
        for p in list(first.paragraphs):
            p._element.getparent().remove(p._element)

        footer = section.footer
        for p in list(footer.paragraphs):
            p._element.getparent().remove(p._element)

        p = footer.add_paragraph()
        p.alignment = AL.CENTER
        pf = p.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.left_indent = Cm(0)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0

        _field_paragraph(p, ' PAGE ', '1')


# ==================== ОГЛАВЛЕНИЕ ====================

def insert_toc(doc, after_p):
    new_p = OxmlElement('w:p')
    after_p._element.addnext(new_p)
    p = Paragraph(new_p, after_p._parent)
    p.style = doc.styles['Normal']
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.left_indent = Cm(0)
    pf.line_spacing = 1.5

    def add(child, size=14):
        r = OxmlElement('w:r')
        _format_run(r, size=size)
        r.append(child)
        p._element.append(r)

    def fld(kind):
        el = OxmlElement('w:fldChar')
        el.set(qn('w:fldCharType'), kind)
        return el

    add(fld('begin'))

    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    add(instr)

    add(fld('separate'))

    r = OxmlElement('w:r')
    _format_run(r, size=14)
    t = OxmlElement('w:t')
    t.text = 'Оглавление не обновлено. Нажмите Ctrl+A, затем F9.'
    r.append(t)
    p._element.append(r)

    add(fld('end'))


def enable_update_fields(doc):
    settings = doc.settings.element
    _drop(settings, 'updateFields')
    el = OxmlElement('w:updateFields')
    el.set(qn('w:val'), 'true')
    settings.append(el)


# ==================== ПРИЛОЖЕНИЯ ====================

def normalize_appendix_title(text):
    up = text.upper().strip()

    if re.match(r'^ПРИЛОЖЕНИЕ\s+[А-Я]', up):
        return up

    m = re.match(r'^ПРИЛОЖЕНИЕ\s*(\d+)', up)
    if m:
        num = int(m.group(1))
        if 1 <= num <= 26:
            letter = chr(ord('А') + num - 1)
            rest = up[m.end():].strip()
            return f'ПРИЛОЖЕНИЕ {letter}' + (f' {rest}' if rest else '')

    if up == 'ПРИЛОЖЕНИЕ':
        return 'ПРИЛОЖЕНИЕ А'

    return up


# ==================== ДВУХЭТАПНАЯ ОБРАБОТКА ====================

def format_to_gost(input_path, output_path):
    """Этап 1: оформление текста БЕЗ разбиения таблиц."""
    doc = Document(input_path)

    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(1.5)

    setup_styles(doc)
    hs = {lvl: _heading_style(doc, lvl) for lvl in (1, 2, 3)}
    by_id = {s._element.get(qn('w:styleId')): s for s in doc.styles}

    add_page_numbers(doc)

    paragraphs = list(doc.paragraphs)
    start = 0
    for i, p in enumerate(paragraphs):
        if p.text.strip().upper() in ('СОДЕРЖАНИЕ', 'ВВЕДЕНИЕ', 'ОГЛАВЛЕНИЕ'):
            start = i
            break

    in_bib = False
    in_contents = False
    toc_p = None
    junk = []

    for i, p in enumerate(paragraphs):
        if i < start:
            continue

        text = p.text.strip()
        up = text.upper()

        if in_contents:
            has_page_break = any(
                br.get(qn('w:type')) == 'page'
                for br in p._element.iter(qn('w:br'))
            )
            ends = (has_page_break or up in ('ВВЕДЕНИЕ', 'ЗАКЛЮЧЕНИЕ')
                    or up.startswith('СПИСОК'))
            if ends:
                in_contents = False
            else:
                has_sect = p._element.find(qn('w:pPr')) is not None and \
                           p._element.find(qn('w:pPr')).find(qn('w:sectPr')) is not None
                if not has_sect:
                    junk.append(p)
                    continue

        ptype = detect(p, in_bib, by_id)

        if ptype == 'toc_entry':
            junk.append(p)
            continue

        if ptype == 'contents_title':
            clear_p(p)
            p.style = doc.styles['Normal']
            replace_text(p, 'СОДЕРЖАНИЕ')
            style_runs(p, size=14, bold=True)
            fmt_p(p, align=AL.CENTER, indent=0, line=2.0)
            toc_p = p
            in_contents = True
            continue

        clear_p(p)

        if ptype == 'bib_title':
            in_bib = True
            replace_text(p, 'СПИСОК ИСТОЧНИКОВ')
            apply_heading(p, 1, hs)
        elif ptype == 'appendix':
            new_text = normalize_appendix_title(p.text)
            replace_text(p, new_text)
            apply_heading(p, 1, hs)
        elif ptype == 'heading1':
            in_bib = False
            apply_heading(p, 1, hs)
        elif ptype in ('heading2', 'heading3'):
            apply_heading(p, int(ptype[-1]), hs)
        elif ptype == 'bib_item':
            apply_body(p, doc)
        elif ptype == 'fig_caption':
            apply_body(p, doc, size=12, align=AL.CENTER, indent=0)
            insert_blank_line_after(p)
        elif ptype == 'table_caption':
            apply_body(p, doc, size=14, align=AL.LEFT, indent=0)
        else:
            apply_body(p, doc)

        print(f'[{ptype:16}] {text[:60]}')

    for p in junk:
        parent = p._element.getparent()
        if parent is not None:
            parent.remove(p._element)

    if toc_p is not None:
        insert_toc(doc, toc_p)

    # Форматируем таблицы, НО НЕ РАЗБИВАЕМ ИХ
    format_tables(doc)
    insert_blank_line_after_tables(doc)

    for p in doc.paragraphs:
        if p.style.name in ('Heading 1', 'Заголовок 1'):
            ensure_page_break(p)

    enable_update_fields(doc)
    doc.save(output_path)


def split_tables_only(input_path, output_path):
    """Этап 2: разбиение таблиц по явным разрывам страниц (пользователь ставит их вручную)."""
    doc = Document(input_path)
    split_long_tables(doc)
    format_tables(doc)
    enable_update_fields(doc)
    doc.save(output_path)